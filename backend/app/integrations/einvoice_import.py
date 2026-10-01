"""Uvoz računov za nazaj: Evelope (in drugi programi) — eSLOG 2.0 XML, ZIP ovojnica, UBL/Peppol XML,
Excel (.xlsx) in CSV. Branje je determinističen (brez AI); vsak račun se preveri (neto + DDV = bruto)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from ..tax_engine.money import D, r2
from .bank_import import parse_amount, parse_date


@dataclass
class ParsedInvoice:
    number: str
    issue_date: date
    net: Decimal
    vat: Decimal
    gross: Decimal
    due_date: date | None = None
    service_date: date | None = None
    seller_name: str = ""
    seller_tax: str | None = None
    buyer_name: str = ""
    buyer_tax: str | None = None
    vat_rate: Decimal | None = None
    paid_date: date | None = None
    paid_amount: Decimal | None = None
    credit_note: bool = False
    source: str = ""
    pdf_name: str | None = None
    warnings: list[str] = field(default_factory=list)

    def check(self) -> "ParsedInvoice":
        if abs(self.net + self.vat - self.gross) > Decimal("0.02"):
            self.warnings.append(f"neto {self.net} + DDV {self.vat} ≠ bruto {self.gross}")
        return self

    def to_dict(self) -> dict:
        out = asdict(self)
        for k, v in out.items():
            if isinstance(v, Decimal):
                out[k] = str(v)
            elif isinstance(v, date):
                out[k] = v.isoformat()
        return out


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _children(el, name):
    return [c for c in el if _local(c.tag) == name]


def _find(el, path: str):
    """Pot brez imenskih prostorov, npr. 'S_BGM/C_C106/D_1004'."""
    cur = [el]
    for part in path.split("/"):
        nxt = []
        for c in cur:
            nxt += _children(c, part)
        cur = nxt
        if not cur:
            return None
    return cur[0]


def _text(el, path: str) -> str:
    x = _find(el, path)
    return (x.text or "").strip() if x is not None else ""


# ------------------------------------------------------------------ eSLOG 2.0
def parse_eslog(root, source: str) -> list[ParsedInvoice]:
    out = []
    for m in root.iter():
        if _local(m.tag) != "M_INVOIC":
            continue
        number = _text(m, "S_BGM/C_C106/D_1004")
        doc_type = _text(m, "S_BGM/C_C002/D_1001")
        dates = {}
        for dtm in _children(m, "S_DTM"):
            dates[_text(dtm, "C_C507/D_2005")] = _text(dtm, "C_C507/D_2380")[:10]
        due = None
        for sg8 in _children(m, "G_SG8"):
            for dtm in _children(sg8, "S_DTM"):
                if _text(dtm, "C_C507/D_2005") == "13":
                    due = _text(dtm, "C_C507/D_2380")[:10]
        parties = {}
        for sg2 in _children(m, "G_SG2"):
            role = _text(sg2, "S_NAD/D_3035")
            c080 = _find(sg2, "S_NAD/C_C080")
            name = " ".join(x.text.strip() for x in (list(c080) if c080 is not None else []) if x.text and x.text.strip())
            tax = None
            for sg3 in _children(sg2, "G_SG3"):
                q = _text(sg3, "S_RFF/C_C506/D_1153")
                if q in ("VA", "AHP"):
                    tax = tax or _text(sg3, "S_RFF/C_C506/D_1154")
            parties[role] = (name, tax)
        moa = {}
        for sg50 in _children(m, "G_SG50"):
            for s in _children(sg50, "S_MOA"):
                moa[_text(s, "C_C516/D_5025")] = D(_text(s, "C_C516/D_5004") or "0")
        tax_amount, taxable, rates = Decimal(0), Decimal(0), []
        for sg52 in _children(m, "G_SG52"):
            r = _text(sg52, "S_TAX/C_C243/D_5278")
            if r:
                rates.append(D(r) / 100)
            for s in _children(sg52, "S_MOA"):
                code, val = _text(s, "C_C516/D_5025"), D(_text(s, "C_C516/D_5004") or "0")
                if code == "124":
                    tax_amount += val
                elif code == "125":
                    taxable += val
        net = moa.get("389", moa.get("79", taxable))
        vat = moa.get("176", tax_amount)
        gross = moa.get("388", moa.get("9", net + vat))
        seller = parties.get("SE") or parties.get("II") or ("", None)
        buyer = parties.get("BY") or parties.get("IV") or ("", None)
        credit = doc_type == "381"
        sign = Decimal(-1) if credit and net > 0 else Decimal(1)
        issue = dates.get("137") or dates.get("3") or next(iter(dates.values()), None)
        if not number or not issue:
            continue
        out.append(ParsedInvoice(
            number=number, issue_date=date.fromisoformat(issue), net=r2(net * sign), vat=r2(vat * sign),
            gross=r2(gross * sign), due_date=date.fromisoformat(due) if due else None,
            service_date=date.fromisoformat(dates["35"]) if dates.get("35") else None,
            seller_name=seller[0], seller_tax=seller[1], buyer_name=buyer[0], buyer_tax=buyer[1],
            vat_rate=max(rates) if rates else None, credit_note=credit, source=source,
        ).check())
    return out


# ------------------------------------------------------------------ UBL 2.1 / Peppol
def parse_ubl(root, source: str) -> list[ParsedInvoice]:
    def t(path):
        return _text(root, path)

    def party(kind):
        p = _find(root, f"{kind}/Party")
        if p is None:
            return "", None
        name = _text(p, "PartyLegalEntity/RegistrationName") or _text(p, "PartyName/Name")
        tax = _text(p, "PartyTaxScheme/CompanyID") or None
        return name, tax
    credit = _local(root.tag) == "CreditNote"
    lmt = "LegalMonetaryTotal"
    net = D(t(f"{lmt}/TaxExclusiveAmount") or t(f"{lmt}/LineExtensionAmount") or "0")
    gross = D(t(f"{lmt}/TaxInclusiveAmount") or t(f"{lmt}/PayableAmount") or "0")
    vat = D(t("TaxTotal/TaxAmount") or str(gross - net))
    rate = t("TaxTotal/TaxSubtotal/TaxCategory/Percent")
    sign = Decimal(-1) if credit else Decimal(1)
    s, b = party("AccountingSupplierParty"), party("AccountingCustomerParty")
    return [ParsedInvoice(
        number=t("ID"), issue_date=date.fromisoformat(t("IssueDate")), net=r2(net * sign), vat=r2(vat * sign),
        gross=r2(gross * sign), due_date=date.fromisoformat(t("DueDate")) if t("DueDate") else None,
        seller_name=s[0], seller_tax=s[1], buyer_name=b[0], buyer_tax=b[1],
        vat_rate=D(rate) / 100 if rate else None, credit_note=credit, source=source).check()]


def parse_xml(data: bytes, source: str) -> list[ParsedInvoice]:
    root = ET.fromstring(data)
    if any(_local(e.tag) == "M_INVOIC" for e in root.iter()):
        return parse_eslog(root, source)
    if _local(root.tag) in ("Invoice", "CreditNote") and "ubl" in root.tag.lower():
        return parse_ubl(root, source)
    return []


# ------------------------------------------------------------------ Excel / CSV (npr. Evelope izvoz seznama računov)
COLS = {
    "number": ["številka računa", "stevilka racuna", "št. računa", "st. racuna", "številka", "stevilka", "račun", "racun", "invoice number", "number"],
    "issue_date": ["datum izdaje", "datum računa", "datum racuna", "datum", "izdano", "issue date", "date"],
    "due_date": ["rok plačila", "rok placila", "zapadlost", "datum zapadlosti", "valuta", "due date"],
    "service_date": ["datum opravljene storitve", "datum storitve", "datum dobave", "opravljeno"],
    "buyer_name": ["kupec", "stranka", "partner", "naziv kupca", "prejemnik", "customer", "client"],
    "buyer_tax": ["id za ddv", "davčna številka", "davcna stevilka", "id ddv kupca", "vat id"],
    "net": ["znesek brez ddv", "osnova", "neto", "brez ddv", "vrednost brez ddv", "net", "subtotal"],
    "vat": ["ddv", "znesek ddv", "vat", "tax"],
    "gross": ["skupaj", "za plačilo", "za placilo", "znesek z ddv", "bruto", "z ddv", "total", "znesek"],
    "paid_date": ["datum plačila", "datum placila", "plačano dne", "placano dne", "paid date"],
    "paid_amount": ["plačano", "placano", "plačan znesek", "paid"],
}


def _map(header: list[str]) -> dict[str, int]:
    norm = [str(h or "").strip().lower() for h in header]
    out = {}
    for key, names in COLS.items():
        for name in names:
            if name in norm and norm.index(name) not in out.values():
                out[key] = norm.index(name)
                break
    return out


def _cell_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return parse_date(str(v))


def _cell_amount(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return r2(D(v))
    return r2(parse_amount(str(v)))


def parse_table(rows: list[list], source: str) -> tuple[list[ParsedInvoice], list[str]]:
    hdr_i, mapping = None, {}
    for i, row in enumerate(rows[:20]):
        m = _map(row)
        if "number" in m and "issue_date" in m and ({"net", "gross"} & set(m)):
            hdr_i, mapping = i, m
            break
    if hdr_i is None:
        raise ValueError("V tabeli ni prepoznane glave (potrebni: številka računa, datum, znesek).")
    out, warns = [], []
    for row in rows[hdr_i + 1:]:
        if not row or all(c in (None, "") for c in row):
            continue
        get = lambda k: row[mapping[k]] if k in mapping and mapping[k] < len(row) else None  # noqa: E731
        try:
            number = str(get("number") or "").strip()
            if not number or number.lower().startswith(("skupaj", "total")):
                continue
            net, vat, gross = _cell_amount(get("net")), _cell_amount(get("vat")), _cell_amount(get("gross"))
            if vat is None and net is not None and gross is not None:
                vat = gross - net
            if net is None and gross is not None:
                net = gross - (vat or Decimal(0))
            if gross is None and net is not None:
                gross = net + (vat or Decimal(0))
            paid_amount = get("paid_amount")
            paid_date = _cell_date(get("paid_date"))
            pa = None
            if isinstance(paid_amount, str) and paid_amount.strip().lower() in ("da", "plačano", "placano", "yes", "paid"):
                pa = gross
            elif paid_amount not in (None, "") and not isinstance(paid_amount, str):
                pa = _cell_amount(paid_amount)
            elif isinstance(paid_amount, str) and re.search(r"\d", paid_amount):
                pa = _cell_amount(paid_amount)
            if paid_date and pa is None:
                pa = gross
            out.append(ParsedInvoice(number=number, issue_date=_cell_date(get("issue_date")), net=r2(net), vat=r2(vat or 0),
                                     gross=r2(gross), due_date=_cell_date(get("due_date")),
                                     service_date=_cell_date(get("service_date")),
                                     buyer_name=str(get("buyer_name") or "").strip(),
                                     buyer_tax=str(get("buyer_tax") or "").strip() or None,
                                     paid_date=paid_date, paid_amount=pa, source=source).check())
        except (ValueError, InvalidOperation, TypeError) as e:
            warns.append(f"Vrstica preskočena ({e}): {[c for c in row][:6]}")
    return out, warns


def parse_xlsx(data: bytes, source: str):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    return parse_table(rows, source)


def parse_csv_table(data: bytes, source: str):
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    delim = ";" if text[:2000].count(";") >= text[:2000].count(",") else ","
    return parse_table(list(csv.reader(io.StringIO(text), delimiter=delim)), source)


# ------------------------------------------------------------------ vstopna točka
@dataclass
class ImportBundle:
    invoices: list[ParsedInvoice] = field(default_factory=list)
    pdfs: dict[str, bytes] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    formats: set[str] = field(default_factory=set)


def parse_file(data: bytes, filename: str) -> ImportBundle:
    b = ImportBundle()
    name = filename.lower()
    if name.endswith(".zip") or data[:2] == b"PK" and not name.endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if info.is_dir() or info.file_size > 30 * 1024 * 1024:
                    continue
                inner = z.read(info.filename)
                if info.filename.lower().endswith(".pdf"):
                    b.pdfs[info.filename] = inner
                elif info.filename.lower().endswith((".xml", ".zip", ".xlsx", ".csv")):
                    sub = parse_file(inner, info.filename)
                    b.invoices += sub.invoices
                    b.pdfs.update(sub.pdfs)
                    b.warnings += sub.warnings
                    b.formats |= sub.formats
        _attach_pdfs(b)
        return b
    if name.endswith(".xml") or data.lstrip()[:1] == b"<":
        try:
            found = parse_xml(data, filename)
        except ET.ParseError as e:
            b.warnings.append(f"{filename}: neveljaven XML ({e})")
            return b
        if not found:
            # ovojnica e-računa ali drug XML brez računa — ni napaka
            if not re.search(rb"envelope|ovojnica|package", data[:500], re.I):
                b.warnings.append(f"{filename}: v XML ni prepoznanega računa (podprta sta eSLOG 2.0 in UBL).")
        b.invoices += found
        b.formats.add("eSLOG 2.0" if b"M_INVOIC" in data else "UBL")
        return b
    if name.endswith(".xlsx"):
        inv, w = parse_xlsx(data, filename)
        b.invoices += inv
        b.warnings += w
        b.formats.add("Excel")
        return b
    inv, w = parse_csv_table(data, filename)
    b.invoices += inv
    b.warnings += w
    b.formats.add("CSV")
    return b


def _attach_pdfs(b: ImportBundle):
    """PDF iz ZIP ovojnice pripne k računu z istim imenom ali številko v imenu datoteke."""
    for inv in b.invoices:
        stem = re.sub(r"\.[^.]+$", "", inv.source.split("/")[-1]).lower()
        token = re.sub(r"[^0-9a-z]", "", inv.number.lower())
        for pdf in b.pdfs:
            p = re.sub(r"\.[^.]+$", "", pdf.split("/")[-1]).lower()
            if p == stem or (token and token in re.sub(r"[^0-9a-z]", "", p)):
                inv.pdf_name = pdf
                break
